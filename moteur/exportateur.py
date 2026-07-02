#!/usr/bin/env python3
"""
Export multi-format (PDF, TXT, MD, DOCX) et fusion PDF.
"""

from __future__ import annotations

import os
import re
import json
from typing import Any

from xml.sax.saxutils import escape as xml_escape
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable

from . import config as cfg
from . import editeur as edt
from .utils import log_fixe

# Imports optionnels
_HAS_DOCX: bool = False
try:
    from docx import Document
    _HAS_DOCX = True
except ImportError:
    pass


# ─────────────────────────────────────────────
#   ROUTEUR D'EXPORT
# ─────────────────────────────────────────────

def _injecter_edits(
    pages_traduites_dict: dict[str, str],
    chemin_sortie: str,
    nom_base: str,
) -> dict[str, str]:
    """Injecte les éditions utilisateur dans les pages traduites si elles existent."""
    conf = cfg.get_config()
    edits = edt.charger_edits(nom_base, os.path.dirname(chemin_sortie))
    if edits is not None:
        # Les éditions sont au format {page_idx: texte_modifie}
        try:
            edits_par_page = json.loads(edits)
            if isinstance(edits_par_page, dict):
                for k, v in edits_par_page.items():
                    if k in pages_traduites_dict and isinstance(v, str) and v.strip():
                        pages_traduites_dict[k] = v
            log_fixe(f" \U0001f4dd \u00c9ditions inject\u00e9es pour {nom_base}")
        except (json.JSONDecodeError, TypeError):
            # Format simple : texte complet → on l'utilise comme page 0
            if isinstance(edits, str) and edits.strip():
                pages_traduites_dict["0"] = edits
                log_fixe(f" \U0001f4dd Texte \u00e9dit\u00e9 inject\u00e9 pour {nom_base}")
    return pages_traduites_dict


def exporter_document(
    pages_traduites_dict: dict[str, str],
    total_pages: int,
    chemin_sortie: str,
    police: str,
    police_bold: str,
    nom_base: str = "",
) -> str:
    """Exporte les pages traduites dans le format configuré."""
    conf = cfg.get_config()
    os.makedirs(os.path.dirname(chemin_sortie), exist_ok=True)

    # Injection des éditions utilisateur
    if nom_base:
        pages_traduites_dict = _injecter_edits(pages_traduites_dict, chemin_sortie, nom_base)

    fmt = conf.format_export.lower()
    if fmt == "txt":
        return _exporter_txt(pages_traduites_dict, total_pages, chemin_sortie)
    elif fmt == "md":
        return _exporter_md(pages_traduites_dict, total_pages, chemin_sortie)
    elif fmt == "docx":
        return _exporter_docx(pages_traduites_dict, total_pages, chemin_sortie)
    else:
        return _exporter_pdf(pages_traduites_dict, total_pages, chemin_sortie, police, police_bold)


# ─────────────────────────────────────────────
#   PDF (reportlab)
# ─────────────────────────────────────────────

def _get_page_size():
    s = cfg.get_config().styles_pdf.get("page_size", "A4")
    return letter if s == "letter" else A4


def _exporter_pdf(
    pages_traduites_dict: dict[str, str],
    total_pages: int,
    chemin_sortie: str,
    police: str,
    police_bold: str,
) -> str:
    """Génère un PDF formaté à partir des pages traduites."""
    s = cfg.get_config().styles_pdf
    m = lambda v: v * cm

    doc = SimpleDocTemplate(
        chemin_sortie,
        pagesize=_get_page_size(),
        rightMargin=m(float(s.get("margin_right", 2.5))),
        leftMargin=m(float(s.get("margin_left", 2.5))),
        topMargin=m(float(s.get("margin_top", 2.5))),
        bottomMargin=m(float(s.get("margin_bottom", 2.5))),
    )
    font_size = float(s.get("font_size", 11))
    line_spacing = float(s.get("line_spacing", 17))
    font_color = colors.HexColor(s.get("font_color", "#222222"))
    chapter_size = float(s.get("chapter_font_size", 13))
    chapter_color = colors.HexColor(s.get("chapter_color", "#222222"))

    style_normal = ParagraphStyle(
        "N", fontName=police, fontSize=font_size,
        leading=line_spacing, spaceAfter=6, textColor=font_color,
    )
    style_chapitre = ParagraphStyle(
        "C", fontName=police_bold, fontSize=chapter_size,
        leading=chapter_size + 7, spaceAfter=8, spaceBefore=14,
        textColor=chapter_color,
    )

    contenu: list[Any] = []
    for i in range(total_pages):
        texte = pages_traduites_dict.get(str(i), "")
        if not texte.strip():
            continue
        for para in re.split(r'\n{2,}', texte):
            para = para.replace('\n', ' ').strip()
            if not para:
                continue
            para_xml = xml_escape(para)
            try:
                if re.match(r'^(Chapitre|Chapter)\s+\d+', para):
                    contenu.append(Paragraph(para_xml, style_chapitre))
                    contenu.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cccccc")))
                    contenu.append(Spacer(1, 0.3 * cm))
                else:
                    contenu.append(Paragraph(para_xml, style_normal))
                    contenu.append(Spacer(1, 0.1 * cm))
            except Exception:
                pass
    doc.build(contenu)
    return chemin_sortie


# ─────────────────────────────────────────────
#   TXT
# ─────────────────────────────────────────────

def _exporter_txt(
    pages_traduites_dict: dict[str, str],
    total_pages: int,
    chemin_sortie: str,
) -> str:
    """Exporte en texte brut."""
    chemin = chemin_sortie.replace(".pdf", ".txt")
    lignes: list[str] = []
    for i in range(total_pages):
        texte = pages_traduites_dict.get(str(i), "")
        if texte.strip():
            lignes.append(texte)
    with open(chemin, "w", encoding="utf-8") as f:
        f.write("\n\n".join(lignes))
    return chemin


# ─────────────────────────────────────────────
#   MARKDOWN
# ─────────────────────────────────────────────

def _exporter_md(
    pages_traduites_dict: dict[str, str],
    total_pages: int,
    chemin_sortie: str,
) -> str:
    """Exporte en Markdown."""
    chemin = chemin_sortie.replace(".pdf", ".md")
    lignes = ["# Traduction automatique\n"]
    for i in range(total_pages):
        texte = pages_traduites_dict.get(str(i), "")
        if texte.strip():
            for para in re.split(r'\n{2,}', texte):
                p = para.replace('\n', ' ').strip()
                if not p:
                    continue
                if re.match(r'^(Chapitre|Chapter)\s+\d+', p):
                    lignes.append(f"\n## {p}\n")
                else:
                    lignes.append(f"{p}\n\n")
    with open(chemin, "w", encoding="utf-8") as f:
        f.write("".join(lignes))
    return chemin


# ─────────────────────────────────────────────
#   DOCX (optionnel)
# ─────────────────────────────────────────────

def _exporter_docx(
    pages_traduites_dict: dict[str, str],
    total_pages: int,
    chemin_sortie: str,
) -> str:
    """Exporte en DOCX (nécessite python-docx)."""
    if not _HAS_DOCX:
        log_fixe("python-docx non install\u00e9, fallback PDF")
        return _exporter_pdf(
            pages_traduites_dict, total_pages, chemin_sortie,
            "Helvetica", "Helvetica-Bold",
        )
    chemin = chemin_sortie.replace(".pdf", ".docx")
    doc = Document()
    doc.add_heading("Traduction automatique", 0)
    for i in range(total_pages):
        texte = pages_traduites_dict.get(str(i), "")
        if texte.strip():
            for para in re.split(r'\n{2,}', texte):
                p = para.replace('\n', ' ').strip()
                if not p:
                    continue
                if re.match(r'^(Chapitre|Chapter)\s+\d+', p):
                    doc.add_heading(p, 1)
                else:
                    doc.add_paragraph(p)
    doc.save(chemin)
    return chemin


# ─────────────────────────────────────────────
#   FUSION PDF
# ─────────────────────────────────────────────

def fusionner_pdfs(chemins_source: list[str], chemin_sortie: str) -> bool:
    """Fusionne plusieurs PDF en un seul fichier."""
    try:
        from pypdf import PdfWriter, PdfReader
    except ImportError:
        log_fixe("pypdf non install\u00e9 pour la fusion")
        return False
    writer = PdfWriter()
    for chemin in chemins_source:
        if os.path.exists(chemin):
            try:
                reader = PdfReader(chemin)
                for page in reader.pages:
                    writer.add_page(page)
            except Exception:
                pass
    if len(writer.pages) > 0:
        writer.write(chemin_sortie)
        writer.close()
        log_fixe(f"Fusion termin\u00e9e: {len(writer.pages)} pages -> {chemin_sortie}")
        return True
    return False
